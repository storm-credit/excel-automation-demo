from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.validate_excel import validate


def _make_book(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


def test_golden_cases(tmp_path: Path) -> None:
    employees = tmp_path / "employees.xlsx"
    departments = tmp_path / "departments.xlsx"
    output = tmp_path / "result.xlsx"

    _make_book(
        employees,
        ["EmployeeID", "Name", "DepartmentCode", "Email"],
        [
            ["E1", "정상", "D1", "ok@example.com"],
            ["E2", "미등록", "D9", "bad@example.com"],
            ["E3", "누락", "", "missing@example.com"],
            ["E4", "중복1", "D1", "dup1@example.com"],
            ["E4", "중복2", "D1", "dup2@example.com"],
        ],
    )
    _make_book(
        departments,
        ["DepartmentCode", "DepartmentName", "Manager"],
        [["D1", "개발", "김팀장"]],
    )

    summary = validate(employees, departments, output)
    assert summary == {"total": 5, "ok": 1, "error": 4}

    wb = load_workbook(output, data_only=True)
    result = wb["Result"]
    statuses = [result.cell(row=i, column=7).value for i in range(2, 7)]
    assert statuses == ["OK", "UNKNOWN_DEPT", "MISSING_DEPT_CODE", "DUPLICATE_ID", "DUPLICATE_ID"]


def test_required_column_validation(tmp_path: Path) -> None:
    employees = tmp_path / "employees.xlsx"
    departments = tmp_path / "departments.xlsx"
    _make_book(employees, ["Name"], [["홍길동"]])
    _make_book(departments, ["DepartmentCode", "DepartmentName", "Manager"], [["D1", "개발", "김팀장"]])

    try:
        validate(employees, departments, tmp_path / "out.xlsx")
    except ValueError as exc:
        assert "Missing required columns" in str(exc)
    else:
        raise AssertionError("Expected missing-column validation to fail")
