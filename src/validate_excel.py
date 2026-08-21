from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

REQUIRED_EMPLOYEE_COLUMNS = ["EmployeeID", "Name", "DepartmentCode", "Email"]
REQUIRED_DEPARTMENT_COLUMNS = ["DepartmentCode", "DepartmentName", "Manager"]

COLUMN_ALIASES = {
    "EmployeeID": {"EmployeeID", "employee_id", "employee id", "사번", "직원ID", "직원 ID", "직원번호"},
    "Name": {"Name", "이름", "성명", "직원명"},
    "DepartmentCode": {
        "DepartmentCode", "department_code", "department code", "부서코드", "부서 코드", "조직코드", "조직 코드"
    },
    "Email": {"Email", "E-mail", "email_address", "이메일", "메일", "메일주소"},
    "DepartmentName": {"DepartmentName", "department_name", "department name", "부서명", "조직명"},
    "Manager": {"Manager", "관리자", "담당자", "부서장", "조직장"},
}


def _normalize_header(value: object) -> str:
    text = str(value).strip().casefold() if value is not None else ""
    return re.sub(r"[\s_\-./]+", "", text)


_ALIAS_LOOKUP = {
    _normalize_header(alias): canonical
    for canonical, aliases in COLUMN_ALIASES.items()
    for alias in aliases
}


def _canonicalize_headers(raw_headers: list[object], required: Iterable[str]) -> list[str]:
    headers: list[str] = []
    seen_required: set[str] = set()

    for value in raw_headers:
        raw = str(value).strip() if value is not None else ""
        canonical = _ALIAS_LOOKUP.get(_normalize_header(raw), raw)
        if canonical in required:
            if canonical in seen_required:
                raise ValueError(f"Duplicate mapped column: {canonical}")
            seen_required.add(canonical)
        headers.append(canonical)

    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")
    return headers


def _read_rows(path: Path, required: Iterable[str]) -> list[dict[str, object]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    first_row = next(iterator, None)
    if first_row is None:
        raise ValueError(f"Empty workbook: {path}")

    try:
        headers = _canonicalize_headers(list(first_row), required)
    except ValueError as exc:
        raise ValueError(f"{exc} in {path.name}") from exc

    result: list[dict[str, object]] = []
    for row in iterator:
        if not any(v not in (None, "") for v in row):
            continue
        item = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
        result.append(item)
    return result


def validate(employees_path: Path, departments_path: Path, output_path: Path) -> dict[str, int]:
    employees = _read_rows(employees_path, REQUIRED_EMPLOYEE_COLUMNS)
    departments = _read_rows(departments_path, REQUIRED_DEPARTMENT_COLUMNS)

    dept_map = {
        str(r["DepartmentCode"]).strip(): r
        for r in departments
        if r.get("DepartmentCode") not in (None, "")
    }
    ids = [str(r.get("EmployeeID") or "").strip() for r in employees]
    id_counts = Counter(ids)

    output_rows: list[list[object]] = []
    error_rows: list[list[object]] = []
    counts = Counter()

    for employee in employees:
        employee_id = str(employee.get("EmployeeID") or "").strip()
        name = str(employee.get("Name") or "").strip()
        dept_code = str(employee.get("DepartmentCode") or "").strip()
        email = str(employee.get("Email") or "").strip()
        dept = dept_map.get(dept_code)

        if employee_id and id_counts[employee_id] > 1:
            validation = "DUPLICATE_ID"
            action = "사번 중복 여부 확인 후 1건으로 정리"
        elif not dept_code:
            validation = "MISSING_DEPT_CODE"
            action = "부서코드 입력"
        elif dept is None:
            validation = "UNKNOWN_DEPT"
            action = "부서 마스터에 존재하는 코드로 수정"
        else:
            validation = "OK"
            action = ""

        counts[validation] += 1
        department_name = str(dept.get("DepartmentName") or "") if dept else ""
        manager = str(dept.get("Manager") or "") if dept else ""
        output_rows.append([
            employee_id, name, dept_code, department_name, manager, email, validation
        ])
        if validation != "OK":
            error_rows.append([
                employee_id, name, dept_code, email, validation, action
            ])

    wb = Workbook()
    result_ws = wb.active
    result_ws.title = "Result"
    errors_ws = wb.create_sheet("Errors")
    summary_ws = wb.create_sheet("Summary")

    result_headers = [
        "EmployeeID", "Name", "DepartmentCode", "DepartmentName",
        "Manager", "Email", "Validation"
    ]
    error_headers = [
        "EmployeeID", "Name", "DepartmentCode", "Email", "ErrorType", "RecommendedAction"
    ]

    result_ws.append(result_headers)
    for row in output_rows:
        result_ws.append(row)

    errors_ws.append(error_headers)
    for row in error_rows:
        errors_ws.append(row)

    summary_ws.append(["Metric", "Value"])
    summary_ws.append(["TOTAL", len(output_rows)])
    summary_ws.append(["OK", counts["OK"]])
    summary_ws.append(["ERROR", len(output_rows) - counts["OK"]])
    summary_ws.append(["DUPLICATE_ID", counts["DUPLICATE_ID"]])
    summary_ws.append(["MISSING_DEPT_CODE", counts["MISSING_DEPT_CODE"]])
    summary_ws.append(["UNKNOWN_DEPT", counts["UNKNOWN_DEPT"]])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    ok_fill = PatternFill("solid", fgColor="E2F0D9")
    error_fill = PatternFill("solid", fgColor="FCE4D6")

    for ws in (result_ws, errors_ws, summary_ws):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column_cells in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 40)

    for row in result_ws.iter_rows(min_row=2, min_col=7, max_col=7):
        row[0].fill = ok_fill if row[0].value == "OK" else error_fill
        row[0].font = Font(bold=True)

    for row in errors_ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].fill = error_fill
        row[0].font = Font(bold=True)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {
        "total": len(output_rows),
        "ok": counts["OK"],
        "error": len(output_rows) - counts["OK"],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="직원/부서 Excel 자동 매칭 및 오류 검증")
    parser.add_argument("employees", type=Path)
    parser.add_argument("departments", type=Path)
    parser.add_argument("--output", type=Path, default=Path("result.xlsx"))
    args = parser.parse_args()

    summary = validate(args.employees, args.departments, args.output)
    print(
        f"완료: {args.output} | 전체 {summary['total']}건 | "
        f"정상 {summary['ok']}건 | 오류 {summary['error']}건"
    )


if __name__ == "__main__":
    main()
