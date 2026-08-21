from __future__ import annotations

import argparse
import tempfile
import time
from pathlib import Path

from openpyxl import Workbook

from src.validate_excel import validate


def _make_workbooks(directory: Path, rows: int) -> tuple[Path, Path, Path]:
    employees = directory / "employees-large.xlsx"
    departments = directory / "departments-large.xlsx"
    output = directory / "result-large.xlsx"

    dept_wb = Workbook(write_only=True)
    dept_ws = dept_wb.create_sheet()
    dept_ws.append(["DepartmentCode", "DepartmentName", "Manager"])
    for index in range(1, 101):
        dept_ws.append([f"D{index:03d}", f"부서{index}", f"관리자{index}"])
    dept_wb.save(departments)

    emp_wb = Workbook(write_only=True)
    emp_ws = emp_wb.create_sheet()
    emp_ws.append(["EmployeeID", "Name", "DepartmentCode", "Email"])
    for index in range(1, rows + 1):
        dept = f"D{((index - 1) % 100) + 1:03d}"
        emp_ws.append([f"E{index:06d}", f"직원{index}", dept, f"user{index}@example.com"])
    emp_wb.save(employees)

    return employees, departments, output


def run(rows: int, max_seconds: float) -> float:
    with tempfile.TemporaryDirectory() as temp_dir:
        directory = Path(temp_dir)
        employees, departments, output = _make_workbooks(directory, rows)

        started = time.perf_counter()
        summary = validate(employees, departments, output)
        elapsed = time.perf_counter() - started

        if summary != {"total": rows, "ok": rows, "error": 0}:
            raise RuntimeError(f"Unexpected summary: {summary}")
        if not output.exists():
            raise RuntimeError("Output file was not created")
        if elapsed > max_seconds:
            raise RuntimeError(
                f"Performance smoke failed: {rows:,} rows took {elapsed:.2f}s "
                f"(limit {max_seconds:.2f}s)"
            )

        print(f"PASS: {rows:,} rows validated in {elapsed:.2f}s")
        return elapsed


def main() -> None:
    parser = argparse.ArgumentParser(description="Large Excel validation smoke test")
    parser.add_argument("--rows", type=int, default=10_000)
    parser.add_argument("--max-seconds", type=float, default=20.0)
    args = parser.parse_args()
    run(args.rows, args.max_seconds)


if __name__ == "__main__":
    main()
